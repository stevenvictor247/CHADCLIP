"""Warrior Village (the football side). Every sheet and note here is made up: no real team data in the repo."""

import csv
import importlib
import io
import json
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest
from fastapi.testclient import TestClient

from hq import approvals, boss, db, ledger, llm, relay, worker
from hq.agents.football import ref as ref_mod
from hq.agents.football import scribe as scribe_mod
from hq.agents.football import tells as tells_mod
from hq.agents.football.clock import Clock
from hq.config import FOOTBALL_INBOX, FOOTBALL_OUTBOX, village_of
from hq.football import depth, files, notebook, rulebook, scouting, season, sheet, tagrules, tendencies

HEADER = ["PLAY #", "ODK", "DN", "DIST", "YARD LN", "HASH", "PLAY TYPE", "GN/LS", "RESULT", "PERSONNEL", "OFF FORM",
          "FORMATION TAG", "BACKFIELD", "MOTION", "OFF PLAY"]


def play(n, form="DOUBLES", tag="", back="GUN STRONG", pers=11, kind="Run", dn=1, dist=10, hash_="L", yard=-30, odk="O"):
    return [n, odk, dn, dist, yard, hash_, kind, 3, "", pers, form, tag, back, "", "Zone" if kind == "Run" else "Mesh"]


def write_xlsx(name, rows, header=HEADER):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Made-up playlist"])   # Hudl exports can have a title row above the header
    ws.append(header)
    for r in rows:
        ws.append(r)
    path = FOOTBALL_INBOX / name
    wb.save(path)
    return path


@pytest.fixture(autouse=True)
def _setup(monkeypatch):
    db.init()
    for t in ("football_files", "football_notes", "football_checklist", "approvals", "tasks", "relay_sent"):
        db.execute(f"DELETE FROM {t}")
    for folder in (FOOTBALL_INBOX, FOOTBALL_OUTBOX):
        folder.mkdir(parents=True, exist_ok=True)
        for p in folder.iterdir():
            p.unlink()
    monkeypatch.setattr(depth, "run_if_due", lambda *a: None)       # the Monday jobs depend on the real day of the week
    monkeypatch.setattr(scouting, "run_if_due", lambda *a: None)
    season.CHANGES_PATH.unlink(missing_ok=True)
    rulebook.PATH.unlink(missing_ok=True)
    sent, asked = [], []
    monkeypatch.setattr(relay, "notify", sent.append)
    monkeypatch.setattr(relay, "ask", lambda *a, **k: asked.append(a))
    return {"sent": sent, "asked": asked}


# ---- reading sheets ----

def test_reads_hudl_export_with_title_row_and_numeric_personnel():
    path = write_xlsx("PlaylistData_test.xlsx", [play(1, pers=2), play(2, pers=12)])
    data = sheet.read(path)
    assert [r["personnel"] for r in data["rows"]] == ["02", "12"]
    assert data["rows"][0]["row"] == 3 and sheet.label(data["rows"][0]) == "Play 1 (row 3)"


def test_reads_practice_csv_with_columns_in_any_order(tmp_path):
    p = tmp_path / "practice.csv"
    with p.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["NOTES", "BACKFIELD", "OFF FORM", "PERSONNEL", "FORMATION TAG", "PLAY #", "CONFIDENCE"])
        w.writerow(["", "GUN WEAK", "TREY", "11", "", "7", "High"])
    rows = sheet.read(p)["rows"]
    assert rows[0]["form"] == "TREY" and rows[0]["play"] == "7" and rows[0]["confidence"] == "High"


def test_a_sheet_without_tag_columns_says_so(tmp_path):
    p = tmp_path / "random.csv"
    p.write_text("a,b\n1,2\n")
    with pytest.raises(sheet.SheetError):
        sheet.read(p)


# ---- Ref's rules ----

def flags_for(*rows, header=HEADER):
    path = write_xlsx("rules.xlsx", rows, header)
    data = sheet.read(path)
    return tagrules.check_sheet(data)


def test_clean_plays_have_no_flags():
    r = flags_for(play(1), play(2, form="TRIPS", tag="FSL", back="GUN WEAK"), play(3, form="EMPTY BUNCH", back="EMPTY", pers=10),
                  play(4, form="TREY OFF", back="ACE", pers=12, hash_="R", tag="FIB"))
    assert r["flags"] == [] and r["plays"] == 4


@pytest.mark.parametrize("cells, column, words", [
    ({"form": "DOUBLES STACKED"}, "form", "formation tag"),
    ({"form": "DOUBLES WIDE"}, "form", "never DOUBLES WIDE"),
    ({"back": "PISTOL STRONG"}, "backfield", "Pistol"),
    ({"back": "GUN STRNG"}, "backfield", "GUN STRONG"),
    ({"tag": "CONDESNED"}, "tag", "CONDENSED"),
    ({"tag": "STACKED"}, "tag", "STACK, not STACKED"),
    ({"form": "TRIPS", "tag": "STACK"}, "tag", "Trips"),
    ({"form": "TRIPS", "tag": "SIB"}, "tag", "2x2"),
    ({"tag": "FSL", "hash_": "M"}, "tag", "middle"),
    ({"back": "EMPTY"}, "form", "starts with EMPTY"),
    ({"form": "EMPTY TRIPS"}, "backfield", "BACKFIELD is EMPTY"),
    ({"pers": "1?"}, "personnel", "No \"?\""),
    ({"pers": 34}, "personnel", "more backs"),
])
def test_rule_breaks_are_flagged_to_fix(cells, column, words):
    kw = {"form": "DOUBLES", "tag": "", "back": "GUN STRONG", "pers": 11, "hash_": "L", **cells}
    if kw["form"].startswith("EMPTY") and "back" not in cells:
        kw["back"] = "GUN STRONG"
    r = flags_for(play(1, **kw))
    hits = [f for f in r["flags"] if f["column"] == column and f["level"] == "fix"]
    assert hits and words.lower() in " ".join(f["message"] for f in hits).lower(), r["flags"]


def test_pro_slot_is_only_a_question_and_non_offense_rows_are_skipped():
    penalty = [3, "O", 1, 10, -30, "L", "", -5, "PENALTY", "", "", "", "", "", ""]
    r = flags_for(play(1, form="PRO SLOT"), play(2, odk="D", form="???"), penalty)
    assert [f["level"] for f in r["flags"]] == ["check"] and r["fix"] == 0 and r["untagged"] == []


def test_ranges_compress_play_numbers():
    assert tagrules.ranges(["1", "2", "3", "7", "9", "10"]) == "1-3, 7, 9-10"


# ---- Tells' counting ----

def test_field_zones_and_down_distance():
    assert tendencies.field_zone({"yard": "8"}) == "Low red zone (12 and in)"
    assert tendencies.field_zone({"yard": "30"}) == "High red zone (35 to 21)"
    assert tendencies.field_zone({"yard": "-30"}) is None
    assert tendencies.down_distance({"dn": "3", "dist": "2"}) == "3rd & short (1-3)"
    assert tendencies.down_distance({"dn": "0", "dist": "10"}) is None   # DN 0 = down not tagged


def test_a_strong_tell_is_found_with_its_sample_size():
    rows = [play(i, form="TREY", back="GUN WEAK", kind="Run") for i in range(1, 11)]
    rows += [play(i, form="DOUBLES", kind="Run" if i % 2 else "Pass") for i in range(11, 31)]
    a = tendencies.analyse(sheet.read(write_xlsx("tells.xlsx", rows)))
    assert a["plays"] == 30
    run_tells = a["run_tells"]
    assert run_tells and run_tells[0]["hits"] == 10 and run_tells[0]["plays"] == 10
    text = tendencies.birds_and_rabbits(a, "Testville")
    assert "run 10 of 10 (100%)" in text


def test_no_tell_from_a_tiny_sample():
    rows = [play(i, form="QUADS", kind="Pass") for i in range(1, 4)] + [play(i, kind="Run" if i % 2 else "Pass") for i in range(4, 30)]
    a = tendencies.analyse(sheet.read(write_xlsx("tiny.xlsx", rows)))
    assert not any(t["value"] == "QUADS" for t in a["pass_tells"])   # 3 plays is too few to call it a tell


# ---- the Inbox, Ref and Tells end to end ----

def test_worker_finds_a_new_sheet_and_ref_checks_it_once(_setup):
    write_xlsx("Testville practice.xlsx", [play(1), play(2, form="DOUBLES STACKED")])
    worker.football_tick()
    tasks = db.rows("SELECT agent, params FROM tasks")
    assert [t["agent"] for t in tasks] == ["Ref"]
    worker.football_tick()
    assert len(db.rows("SELECT 1 FROM tasks")) == 1   # not again until the file changes
    out = ref_mod.check_file("Testville practice.xlsx")
    assert "1 to fix" in out and "Ref check.xlsx" in out
    f = files.listing()[0]
    assert f["status"] == "checked" and f["fix"] == 1 and f["ref_copy"]
    copy = openpyxl.load_workbook(FOOTBALL_OUTBOX / f["ref_copy"]).active
    assert any(c.value == "REF" for c in copy[2])
    assert (FOOTBALL_INBOX / "Testville practice.xlsx").exists()   # the original is never changed


def test_ref_reports_an_unreadable_file_instead_of_crashing():
    (FOOTBALL_INBOX / "broken.xlsx").write_bytes(b"not a spreadsheet")
    out = ref_mod.check_file("broken.xlsx")
    assert "couldn't read" in out.lower() or "broken.xlsx" in out
    assert db.rows("SELECT status FROM football_files")[0]["status"] == "failed"


def test_tells_report_waits_for_ok_then_lands_in_outbox(_setup):
    rows = [play(i, form="TREY", kind="Run") for i in range(1, 12)] + [play(i, kind="Pass") for i in range(12, 24)]
    write_xlsx("Davenport week.xlsx", rows)
    out = tells_mod.build("Davenport week.xlsx")
    assert "Davenport" in out
    pending = approvals.pending()
    assert len(pending) == 1 and pending[0]["kind"] == "football_report" and pending[0]["yes_label"] == "Save to Outbox"
    assert _setup["asked"], "the phone got the approval too"
    assert not list(FOOTBALL_OUTBOX.iterdir())
    approvals.decide(pending[0]["id"], "approve")
    saved = list(FOOTBALL_OUTBOX.glob("*.docx"))
    assert len(saved) == 1 and saved[0].name.startswith("Davenport scouting report")


def test_birds_and_rabbits_is_its_own_one_page_document_with_its_own_button(_setup):
    from docx import Document
    rows = [play(i, form="TREY", kind="Run") for i in range(1, 12)] + [play(i, kind="Pass") for i in range(12, 24)]
    write_xlsx("Davenport week.xlsx", rows)
    out = tells_mod.build("Davenport week.xlsx", kind="birds")
    assert out.startswith("Birds and Rabbits for Davenport")
    ap = approvals.pending()[0]
    assert ap["title"].startswith("Birds and Rabbits: Davenport")
    approvals.decide(ap["id"], "approve")
    saved = list(FOOTBALL_OUTBOX.glob("*.docx"))
    assert len(saved) == 1 and saved[0].name.startswith("Davenport birds and rabbits")
    text = "\n".join(p.text for p in Document(saved[0]).paragraphs)
    assert "BIRDS AND RABBITS" in text and "RUN TELLS" in text and "PASS TELLS" in text and "TREY" in text.upper()
    assert "FORMATION" not in text.upper().split("TELLS")[-1]   # no tendency tables: those are in the scouting report
    # the dashboard buttons make the right tasks
    api_client = api()
    r = api_client.post("/api/football/upload", files={"file": ("Ferris week.xlsx", (FOOTBALL_INBOX / "Davenport week.xlsx").read_bytes())},
                        headers={"X-HQ": "1"})
    fid = files.listing()[0]["id"]
    api_client.post(f"/api/football/files/{fid}/report", json={"kind": "birds"}, headers={"X-HQ": "1"})
    api_client.post(f"/api/football/files/{fid}/report", json={}, headers={"X-HQ": "1"})
    params = [json.loads(t["params"]) for t in db.rows("SELECT params FROM tasks WHERE agent = 'Tells' ORDER BY id")]
    assert params[-2].get("kind") == "birds" and "kind" not in params[-1]


def test_tells_refuses_a_sheet_with_too_few_plays():
    write_xlsx("small.xlsx", [play(1), play(2)])
    assert "too few" in tells_mod.build("small.xlsx") and not approvals.pending()


def test_throwing_away_a_report_deletes_the_draft():
    write_xlsx("x.xlsx", [play(i, kind="Run" if i % 3 else "Pass") for i in range(1, 15)])
    tells_mod.build("x.xlsx", "Testville")
    draft = approvals.pending()[0]["payload"]["draft"]
    approvals.decide(approvals.pending()[0]["id"], "reject")
    assert not list(FOOTBALL_OUTBOX.iterdir()) and not Path(draft).exists()


def test_uploads_sort_sheets_rulebook_and_notebook():
    r = files.save_upload("../../evil.xlsx", io.BytesIO(b"x"))
    assert r["kind"] == "sheet" and (FOOTBALL_INBOX / "evil.xlsx").exists()
    with pytest.raises(ValueError):
        files.save_upload("virus.exe", io.BytesIO(b"x"))
    with pytest.raises(ValueError):
        files.save_upload("random.md", io.BytesIO(b"# hi"))
    assert files.save_upload("my rulebook.md", io.BytesIO(("# Rules\n" + "Never invent stats. " * 5).encode()))["kind"] == "rulebook"
    assert rulebook.loaded()
    r = files.save_upload("Coaching Notebook.md", io.BytesIO("## Fronts\n- 9/1/26 — Made-up front note\n".encode()))
    assert r["kind"] == "notebook" and notebook.notes()[0]["day"] == "2026-09-01"


# ---- Clock ----

def test_checklist_states_and_ticking_off():
    g = next(x for x in season.games() if x["opponent"] == "Davenport")
    w = season.week(date(2026, 10, 20))
    assert w["game"]["opponent"] == "Davenport" and w["days_to_game"] == 4
    states = {x["item"]: x["state"] for x in w["checklist"]}
    assert states["Defensive self-scout"] == "late" and states["Personnel sheet"] == "today"
    assert "Teamworks itinerary" not in w["late"]
    season.mark("personnel", True, "Davenport")
    assert {x["item"]: x["state"] for x in season.checklist(g, date(2026, 10, 20))}["Personnel sheet"] == "done"
    season.mark("Personnel sheet", False, "Davenport")
    assert {x["item"]: x["state"] for x in season.checklist(g, date(2026, 10, 20))}["Personnel sheet"] == "today"


def test_nicknames_find_the_opponent():
    assert season.find_opponent("PlaylistData SVSU 10-17") == "Saginaw Valley State"
    assert season.find_opponent("nmu cutups") == "Northern Michigan"
    assert season.find_opponent("nothing here") is None


def test_set_game_and_due_days_are_saved_outside_git():
    season.set_game("Davenport", "2026-10-24", "1:00 PM", "home")
    g = next(x for x in season.games() if x["opponent"] == "Davenport")
    assert g["confirmed"] and g["kickoff"] == "1:00 PM" and season.CHANGES_PATH.exists()
    season.set_due("Personnel sheet", weekday="Wednesday")
    assert next(i for i in season.items() if i["name"] == "Personnel sheet")["days_before"] == 3
    season.set_due("Film cut-ups", 2, add=True)
    assert any(i["name"] == "Film cut-ups" for i in season.items())
    season.set_due("Film cut-ups", remove=True)
    assert not any(i["name"] == "Film cut-ups" for i in season.items())
    with pytest.raises(ValueError):
        season.set_game("Davenport", site="moon")


def test_clock_agent_ticks_items_through_the_boss():
    out = Clock().run({"params": json.dumps({"done": "scouting report"})})
    assert "Scouting report for" in out and "done" in out


def test_morning_reminder_is_sent_once_a_day(_setup):
    assert not season.remind_if_due(8, datetime(2026, 10, 20, 7, 0))   # before 8 am
    assert season.remind_if_due(8, datetime(2026, 10, 20, 9, 0))
    assert not season.remind_if_due(8, datetime(2026, 10, 20, 15, 0))  # already sent today
    assert len(_setup["sent"]) == 1 and "Personnel sheet" in _setup["sent"][0] and "Late:" in _setup["sent"][0]


def test_no_reminder_on_a_quiet_day():
    for i in season.items():   # everything done, mid-week: nothing to say
        season.mark(i["name"], True, "Davenport")
    assert season.reminder(date(2026, 10, 22)) is None


def test_season_over():
    w = season.week(date(2026, 12, 1))
    assert w["season_over"] and season.reminder(date(2026, 12, 1)) is None


# ---- Scribe and the notebook ----

def test_notebook_import_skips_duplicates_and_exports_back():
    md = "# Notebook\n## Fronts\n- 9/14/26 — Made-up front A\n- Made-up front B\n## Formations\n| Name | Meaning |\n|---|---|\n| ZIP | made up |\n"
    assert notebook.import_markdown(md) == 3
    assert notebook.import_markdown(md) == 0
    out = notebook.export_markdown()
    assert "## Fronts" in out and "9/14/26 — Made-up front A" in out and "ZIP: made up" in out


def test_nb_note_is_filed_by_claude(monkeypatch):
    monkeypatch.setattr(llm, "HAS_LLM", True)
    monkeypatch.setattr(ledger, "over_budget", lambda village="business": False)
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"entry": "Made-up blitz is a 5-man pressure.", "tab": "Pressures and blitzes", "conflicts": []})
    out = scribe_mod.note("NB: made up blitz = 5 man pressure")
    assert "Filed under Pressures and blitzes" in out
    assert notebook.notes()[0]["tab"] == "Pressures and blitzes" and not approvals.pending()


def test_a_clash_asks_before_replacing(monkeypatch):
    old = notebook.add("Made-up call means cover 2.", "Coverages")
    monkeypatch.setattr(llm, "HAS_LLM", True)
    monkeypatch.setattr(ledger, "over_budget", lambda village="business": False)
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"entry": "Made-up call now means cover 4.", "tab": "Coverages",
                                                          "conflicts": [{"id": old, "why": "Different coverage."}, {"id": 999, "why": "x"}]})
    out = scribe_mod.note("NB: made-up call is cover 4 now")
    assert "clash" in out and f"#{old}" in out
    a = approvals.pending()[0]
    assert a["kind"] == "football_note" and a["payload"]["old_ids"] == [old]   # unknown ids are dropped
    approvals.decide(a["id"], "approve")
    assert [n["entry"] for n in notebook.notes()] == ["Made-up call now means cover 4."]


def test_without_claude_the_note_is_still_saved(monkeypatch):
    monkeypatch.setattr(llm, "HAS_LLM", False)
    out = scribe_mod.note("NB: something made up")
    assert "Saved as written under General" in out and notebook.notes()[0]["entry"] == "something made up"


def test_boss_sends_nb_messages_straight_to_scribe(monkeypatch):
    monkeypatch.setattr(llm, "HAS_LLM", False)
    reply = boss.chat("NB: made-up note from chat")
    assert "📓" in reply and notebook.notes()[-1]["entry"] == "made-up note from chat"


# ---- budgets and villages ----

def test_football_has_its_own_budget(monkeypatch):
    assert village_of("Ref") == "football" and village_of("Chad") == "business" and village_of("captain") == "football"
    before = ledger.spent_today("football")
    ledger.record("Scribe", "claude-haiku-5-5", SimpleNamespace(input_tokens=1_000_000, output_tokens=0))
    assert ledger.spent_today("football") > before
    monkeypatch.setattr(ledger, "FOOTBALL_DAILY_BUDGET", ledger.spent_today("football") - 0.01)
    assert ledger.over_budget("football")
    monkeypatch.setattr(ledger, "DAILY_BUDGET", 10_000)
    assert not ledger.over_budget()   # the business side keeps going


# ---- the dashboard ----

def api():
    import hq.api
    import hq.config
    for m in (hq.config, db, hq.api):   # other test files reload these against their own temp folders
        importlib.reload(m)
    db.init()
    return TestClient(hq.api.app)


def test_state_has_both_villages():
    s = api().get("/api/state").json()
    villages = {a["name"]: a["village"] for a in s["agents"]}
    assert villages["Chad"] == "business" and villages["Ref"] == "football" and villages["Boss"] == "business"
    assert {"week", "files", "outbox", "notes", "rulebook_loaded", "budget_usd"} <= set(s["football"])


def test_football_posts_need_the_dashboard_header():
    c = api()
    assert c.post("/api/football/note", json={"text": "x"}).status_code == 403
    assert c.post("/api/football/checklist", json={"item": "Personnel sheet"}, headers={"X-HQ": "1"}).status_code == 200


def test_upload_check_and_download_through_the_api():
    c = api()
    wb = openpyxl.Workbook()
    wb.active.append(HEADER)
    wb.active.append(play(1, form="DOUBLES WIDE"))
    buf = io.BytesIO()
    wb.save(buf)
    r = c.post("/api/football/upload", files={"file": ("Testville.xlsx", buf.getvalue())}, headers={"X-HQ": "1"})
    assert r.status_code == 200 and r.json()["kind"] == "sheet"
    assert db.rows("SELECT agent FROM tasks")[0]["agent"] == "Ref"
    ref_mod.check_file("Testville.xlsx")
    f = c.get("/api/state").json()["football"]["files"][0]
    assert f["fix"] == 1 and f["flags"][0]["column"] == "form"
    got = c.get(f"/api/football/outbox/{f['ref_copy']}")
    assert got.status_code == 200 and got.content[:2] == b"PK"
    assert c.get("/api/football/outbox/..%2F..%2Fhq.db").status_code == 404
    assert c.post(f"/api/football/files/{f['id']}/report", json={"opponent": "Davenport"}, headers={"X-HQ": "1"}).status_code == 200


def test_rulebook_save_and_read_through_the_api():
    c = api()
    assert c.post("/api/football/rulebook", json={"text": "short"}, headers={"X-HQ": "1"}).status_code == 400
    text = "# Made-up rulebook\n" + "Tells first, short bullets, sample sizes. " * 3
    assert c.post("/api/football/rulebook", json={"text": text}, headers={"X-HQ": "1"}).status_code == 200
    got = c.get("/api/football/rulebook").json()
    assert got["loaded"] and "Made-up rulebook" in got["text"]


def test_football_approvals_are_labelled_by_village():
    approvals.request("Tells", "football_report", "Made-up report", "preview", {"draft": "x", "name": "y"})
    a = api().get("/api/state").json()["approvals"][0]
    assert a["village"] == "football"


# ---- Captain and the worker ----

def test_captain_tools_do_real_work_and_unknown_tools_fail():
    from hq.agents.football import captain
    assert "done" in captain._tool("mark_done", {"item": "Birds and Rabbits", "opponent": "Davenport"})
    assert "game_week" in captain._tool("football_status", {}) or "season" in captain._tool("football_status", {})
    with pytest.raises(ValueError):
        captain._tool("send_to_coach", {})


def test_captain_without_claude_still_answers(monkeypatch):
    from hq.agents.football import captain
    monkeypatch.setattr(llm, "HAS_LLM", False)
    assert "football status" in captain.ask("what's due?")


def test_football_tasks_ping_the_phone_when_done(_setup):
    write_xlsx("Testville.xlsx", [play(1)])
    db.create_task("Ref", "Check Testville.xlsx", {"file": "Testville.xlsx"}, created_by="test")
    assert worker.run_next_task()
    assert _setup["sent"] and _setup["sent"][-1].startswith("Ref: Testville.xlsx")


def test_a_downloaded_notebook_imports_back_without_extra_entries():
    notebook.add("Made-up entry", "Fronts", day="2026-10-01")
    out = notebook.export_markdown()
    db.execute("DELETE FROM football_notes")
    assert notebook.import_markdown(out) == 1
    assert notebook.shown(notebook.notes()[0]) == "10/1/26 — Made-up entry" and notebook.notes()[0]["tab"] == "Fronts"

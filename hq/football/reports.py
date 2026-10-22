"""The files HQ makes for you: the scouting report (.docx) and Ref's marked-up copy of your sheet (.xlsx).

Report style follows your one-pagers: Arial, WSU green #008168 and gold #c9a227, tells first, short bullets.
"""

from pathlib import Path

from .sheet import TAG_COLUMNS

GREEN, GOLD = "008168", "C9A227"
REF_FILL = {"fix": "F8D7D3", "check": "FBEFC8"}


def safe_name(text: str) -> str:
    keep = "".join(ch if ch.isalnum() or ch in " -_.()&" else "_" for ch in text).strip()
    return " ".join(keep.split())[:120] or "file"


def scouting_docx(a: dict, opponent: str, path: Path) -> Path:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Pt(42)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

    def heading(text, size=13, color=GREEN):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold, r.font.size, r.font.color.rgb, r.font.name = True, Pt(size), RGBColor.from_string(color), "Arial"
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(8), Pt(2)
        return p

    def shade(cell, hex_color):
        tc = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), hex_color)
        tc.append(shd)

    def grid(headers, rows):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style, t.alignment = "Table Grid", WD_TABLE_ALIGNMENT.LEFT
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            c.text = ""
            run = c.paragraphs[0].add_run(h)
            run.bold, run.font.size, run.font.color.rgb = True, Pt(9), RGBColor(255, 255, 255)
            shade(c, GREEN)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)
                for p in cells[i].paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9)

    heading(f"{opponent.upper()} OFFENSE: SCOUTING REPORT", 16)
    p = doc.add_paragraph(f"{a['plays']} plays ({a['run']} run, {a['pass']} pass) from {a['file']}. "
                          "Every number is a count of your tags.")
    p.runs[0].font.size = Pt(8.5)

    heading("BIRDS AND RABBITS", 13, GOLD)
    for title, key, word in (("Run tells", "run_tells", "run"), ("Pass tells", "pass_tells", "pass")):
        heading(title, 11)
        if not a[key]:
            doc.add_paragraph(f"No {word} tell strong enough (needs 5+ plays and a 70%+ lean).", style="List Bullet")
        for t in a[key]:
            doc.add_paragraph(t["text"], style="List Bullet")

    for feature, rows in a["tables"].items():
        if feature == "Motion" and sum(r["plays"] for r in rows) < 5:
            continue
        heading(feature.upper() + (" (TOP 12)" if feature == "Formation" else ""), 11)
        grid([feature, "Plays", "%", "Run", "Pass", "Run %"],
             [(r["value"], r["plays"], f"{r['pct']}%", r["run"], r["pass"], f"{r['run_pct']}%") for r in rows])

    if a["top_plays"]:
        heading("WHAT THEY RUN FROM THEIR TOP FORMATIONS", 11)
        for f in a["top_plays"]:
            calls = ", ".join(f"{c['call']} ({c['n']})" for c in f["calls"])
            doc.add_paragraph(f"{f['formation']} ({f['plays']} plays): {calls}", style="List Bullet")

    notes = []
    if a["skipped"]:
        notes.append(f"{a['skipped']} snaps without a run or pass (penalties, timeouts) are left out.")
    if a["no_down"]:
        notes.append(f"{a['no_down']} plays have no down tagged (DN 0 or blank), so they're not in Down and distance.")
    if notes:
        heading("NOTES", 10)
        for n in notes:
            doc.add_paragraph(n, style="List Bullet")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def ref_copy(source: Path, result: dict, path: Path) -> Path | None:
    """Your sheet with Ref's flags: flagged cells colored (red = fix, yellow = check) and a REF column saying why.
    Only for .xlsx files; your original file is never changed."""
    if source.suffix.lower() not in (".xlsx", ".xlsm"):
        return None
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill

    from .sheet import COLUMNS, _find_header, _norm_header
    wb = openpyxl.load_workbook(source)
    ws = None
    for sheet in wb.worksheets:
        grid = [[c.value for c in row] for row in sheet.iter_rows(max_row=15)]
        if _find_header(grid) is not None:
            ws, header_row = sheet, _find_header(grid) + 1
            break
    if ws is None:
        return None
    where = {}
    for cell in ws[header_row]:
        name = _norm_header(cell.value)
        for ours, spellings in COLUMNS.items():
            if name in spellings and ours not in where:
                where[ours] = cell.column
    ref_col = ws.max_column + 1
    ws.cell(header_row, ref_col, "REF").font = Font(bold=True)
    by_row = {}
    for f in result["flags"]:
        by_row.setdefault(f["row"], []).append(f)
    for r, flags in by_row.items():
        notes = []
        for f in flags:
            col = where.get(f["column"])
            if col:
                cell = ws.cell(r, col)
                worst = "fix" if any(x["level"] == "fix" and x["column"] == f["column"] for x in flags) else "check"
                cell.fill = PatternFill("solid", fgColor=REF_FILL[worst])
                cell.comment = Comment(f["message"], "Ref")
            notes.append(f"{f['level'].upper()}: {f['message']}")
        ws.cell(r, ref_col, " | ".join(notes))
    ws.column_dimensions[openpyxl.utils.get_column_letter(ref_col)].width = 70
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


__all__ = ["scouting_docx", "ref_copy", "safe_name", "TAG_COLUMNS"]
